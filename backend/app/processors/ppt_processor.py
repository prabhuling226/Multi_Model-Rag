from pathlib import Path

from pptx import Presentation

from app.processors.base_processor import BaseProcessor
from app.processors.pdf.document_element import DocumentElement

class PPTProcessor(BaseProcessor):

    async def process(
        self,
        file_path: Path
    ) -> dict:

        prs = Presentation(file_path)

        text = ""

        for slide in prs.slides:

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text += shape.text + "\n"

        element = DocumentElement(
            element_type="TEXT",
            page_number=1,
            content=text,
            metadata={"slides": len(prs.slides)}
        )

        return {
            "elements": [element]
        }